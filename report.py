"""Parse a Jira export (CSV or XLSX) and compute progress metrics.

No persistence: the file is read in memory, summarized, and discarded.
"""
import csv
import io
import re

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

COMPLETE_STATUSES = ["QA Passed", "Deployed To Production"]
INCOMPLETE_STATUSES = ["To Do", "In Progress", "Dev to Test", "REVISION"]

# Default column mapping. Each value can be a header name or a column letter (e.g. "J").
POINTS_COLUMN = "Custom field (Story point estimate)"
STATUS_COLUMN = "Status"

KEY_HEADERS = ["issue key", "key"]
SUMMARY_HEADERS = ["summary"]


class ReportError(Exception):
    pass


def _norm(value):
    return str(value or "").replace("﻿", "").strip().lower()


def _read_rows(filename, data):
    name = filename.lower()
    if name.endswith(".csv"):
        text = data.decode("utf-8-sig", errors="replace")
        return [row for row in csv.reader(io.StringIO(text))]
    if name.endswith((".xlsx", ".xlsm")):
        wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        ws = wb.active
        rows = [["" if c is None else c for c in row] for row in ws.iter_rows(values_only=True)]
        wb.close()
        return rows
    raise ReportError("Unsupported file type. Upload a .csv or .xlsx Jira export.")


def _find_header_row(rows, wanted):
    """Jira exports put headers on row 1, but tolerate a few title rows above."""
    wanted = {_norm(w) for w in wanted}
    for i, row in enumerate(rows[:10]):
        if any(_norm(c) in wanted for c in row):
            return i
    return 0


def _letter_to_index(text):
    """'J' -> 9, 'AA' -> 26. Returns None if text isn't a column letter."""
    t = text.strip().upper()
    if not re.fullmatch(r"[A-Z]{1,3}", t):
        return None
    n = 0
    for ch in t:
        n = n * 26 + (ord(ch) - 64)
    return n - 1


def _resolve(header, spec, label):
    """Find a column by header name first, then by column letter. Error if neither works."""
    normed = [_norm(h) for h in header]
    if _norm(spec) in normed:
        return normed.index(_norm(spec))
    idx = _letter_to_index(spec)
    if idx is not None and idx < len(header):
        return idx
    found = ", ".join(f"{get_column_letter(i + 1)}: {h}" for i, h in enumerate(header) if str(h).strip())
    raise ReportError(
        f"Couldn't find the {label} column “{spec}” in this file. "
        f"Enter a header name or column letter under Column mapping. Columns in this file: {found}"
    )


def _optional_col(header, candidates):
    normed = [_norm(h) for h in header]
    for cand in candidates:
        if cand in normed:
            return normed.index(cand)
    return None


def _to_points(value):
    """Returns (points or None, is_invalid)."""
    if value is None or str(value).strip() == "":
        return None, False
    try:
        return float(value), False
    except (TypeError, ValueError):
        return None, True


def build_report(filename, data, complete=None, incomplete=None,
                 points_col=None, status_col=None):
    complete = complete or COMPLETE_STATUSES
    incomplete = incomplete or INCOMPLETE_STATUSES
    points_col = (points_col or POINTS_COLUMN).strip()
    status_col = (status_col or STATUS_COLUMN).strip()
    complete_set = {s.strip().lower() for s in complete}
    incomplete_set = {s.strip().lower() for s in incomplete}

    rows = _read_rows(filename, data)
    if not rows:
        raise ReportError("The file is empty.")

    h = _find_header_row(rows, [status_col, points_col, "status"])
    header = rows[h]
    pts_i = _resolve(header, points_col, "story points")
    st_i = _resolve(header, status_col, "status")
    if pts_i == st_i:
        raise ReportError("Story points and status are mapped to the same column. Check Column mapping.")
    key_i = _optional_col(header, KEY_HEADERS)
    sum_i = _optional_col(header, SUMMARY_HEADERS)

    totals = {
        "points_complete": 0.0, "points_incomplete": 0.0,
        "tasks_complete": 0, "tasks_incomplete": 0,
    }
    by_status = {}
    unestimated = []   # counted as tasks, 0 points
    unmapped = []      # status not in either list; excluded from totals
    invalid_points = []  # non-numeric value in the points column

    for row in rows[h + 1:]:
        if not any(str(c).strip() for c in row):
            continue
        get = lambda i: row[i] if i is not None and i < len(row) else ""
        status = str(get(st_i)).strip()
        raw_points = get(pts_i)
        points, bad = _to_points(raw_points)
        item = {"key": str(get(key_i)), "summary": str(get(sum_i)), "status": status,
                "value": str(raw_points)}

        s = status.lower()
        if s in complete_set:
            bucket = "complete"
        elif s in incomplete_set:
            bucket = "incomplete"
        else:
            unmapped.append(item)
            continue

        totals[f"tasks_{bucket}"] += 1
        totals[f"points_{bucket}"] += points or 0
        if bad:
            invalid_points.append(item)
        elif points is None:
            unestimated.append(item)

        entry = by_status.setdefault(status, {"bucket": bucket, "tasks": 0, "points": 0.0})
        entry["tasks"] += 1
        entry["points"] += points or 0

    total_pts = totals["points_complete"] + totals["points_incomplete"]
    total_tasks = totals["tasks_complete"] + totals["tasks_incomplete"]

    if total_tasks == 0 and unmapped:
        sample = ", ".join(sorted({i["status"] or "(blank)" for i in unmapped})[:5])
        raise ReportError(
            f"No rows matched your statuses. The status column is reading "
            f"{get_column_letter(st_i + 1)} · “{header[st_i]}”, which contains values like: {sample}. "
            f"Check Column mapping (or Status mapping)."
        )
    numeric = total_tasks - len(invalid_points) - len(unestimated)
    if invalid_points and numeric == 0:
        raise ReportError(
            f"The story points column is reading {get_column_letter(pts_i + 1)} · “{header[pts_i]}”, "
            f"which has no numbers in it. Check Column mapping."
        )
    totals["points_total"] = total_pts
    totals["tasks_total"] = total_tasks
    totals["points_pct"] = round(100 * totals["points_complete"] / total_pts, 1) if total_pts else 0
    totals["tasks_pct"] = round(100 * totals["tasks_complete"] / total_tasks, 1) if total_tasks else 0

    return {
        "filename": filename,
        "totals": totals,
        "by_status": sorted(by_status.items(), key=lambda kv: (kv[1]["bucket"] != "complete", kv[0])),
        "unestimated": unestimated,
        "unmapped": unmapped,
        "invalid_points": invalid_points,
        "columns": {
            "points": f"{get_column_letter(pts_i + 1)} · {header[pts_i]}",
            "status": f"{get_column_letter(st_i + 1)} · {header[st_i]}",
        },
        "complete": complete,
        "incomplete": incomplete,
    }


def fmt(n):
    """2.0 -> '2', 2.5 -> '2.5'."""
    return f"{n:g}"
